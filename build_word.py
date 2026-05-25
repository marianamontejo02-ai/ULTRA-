from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin   = Cm(2.2)
    section.right_margin  = Cm(2.2)

# ── Color palette ─────────────────────────────────────────────────────────────
LIME    = RGBColor(0x80, 0xCC, 0x00)   # lime (darker for print readability)
DARK    = RGBColor(0x1C, 0x1C, 0x1C)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY   = RGBColor(0x66, 0x66, 0x66)
MGRAY   = RGBColor(0x99, 0x99, 0x99)
LIME_BG = RGBColor(0xF0, 0xFF, 0xCC)   # very light lime for cell bg

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_font(run, name='Calibri', size=11, bold=False,
             color=None, italic=False):
    run.font.name  = name
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def cell_bg(cell, hex_color):
    """Set table cell background color."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def cell_borders(cell, color='AAAAAA', sz=4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'single')
        el.set(qn('w:sz'),    str(sz))
        el.set(qn('w:color'), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)

def no_space(para):
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after  = Pt(0)

def para_color(para, hex6):
    """Shade an entire paragraph background."""
    pPr  = para._p.get_or_add_pPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex6)
    pPr.append(shd)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    no_space(p)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    if level == 1:
        set_font(r, size=15, bold=True, color=DARK)
        # lime underline bar via border
        pPr  = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bot  = OxmlElement('w:bottom')
        bot.set(qn('w:val'),   'single')
        bot.set(qn('w:sz'),    '12')
        bot.set(qn('w:color'), '80CC00')
        pBdr.append(bot)
        pPr.append(pBdr)
    else:
        set_font(r, size=12, bold=True, color=LIME)
    return p

def add_body(doc, text, indent=False, color=None, size=10.5, italic=False):
    p = doc.add_paragraph()
    no_space(p)
    p.paragraph_format.space_after = Pt(3)
    if indent:
        p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(text)
    set_font(r, size=size, color=color or LGRAY, italic=italic)
    return p

def add_bullet(doc, text, bold_prefix=None, indent=0.2):
    p = doc.add_paragraph(style='List Bullet')
    no_space(p)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Inches(indent)
    if bold_prefix:
        rb = p.add_run(bold_prefix + ' ')
        set_font(rb, size=10.5, bold=True, color=DARK)
    r = p.add_run(text)
    set_font(r, size=10.5, color=LGRAY)
    return p

def add_table(doc, headers, rows,
              col_widths=None, header_bg='1C1C1C',
              stripe=True):
    cols  = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hrow = table.rows[0]
    for i, h in enumerate(headers):
        c = hrow.cells[i]
        cell_bg(c, header_bg)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_space(p)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(3)
        r = p.add_run(h)
        set_font(r, size=10, bold=True, color=WHITE)

    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        bg  = 'F5F5F5' if (stripe and ri % 2 == 0) else 'FFFFFF'
        for ci, val in enumerate(row_data):
            c = row.cells[ci]
            cell_bg(c, bg)
            p = c.paragraphs[0]
            no_space(p)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            if ci == 0:
                r = p.add_run(val)
                set_font(r, size=10, bold=True, color=DARK)
            else:
                r = p.add_run(val)
                set_font(r, size=10, color=LGRAY)

    # Col widths
    if col_widths:
        for ri2, row2 in enumerate(table.rows):
            for ci2, cell2 in enumerate(row2.cells):
                cell2.width = Inches(col_widths[ci2])

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def lime_box(doc, title, content_lines):
    """A highlighted lime-bg info box."""
    p_title = doc.add_paragraph()
    no_space(p_title)
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after  = Pt(0)
    para_color(p_title, '80CC00')
    r = p_title.add_run(f'  {title}')
    set_font(r, size=11, bold=True, color=WHITE)

    for line in content_lines:
        pb = doc.add_paragraph()
        no_space(pb)
        pb.paragraph_format.space_after = Pt(1)
        pb.paragraph_format.left_indent = Inches(0.15)
        para_color(pb, 'F0FFCC')
        r2 = pb.add_run(f'  {line}')
        set_font(r2, size=10, color=DARK)

    gap = doc.add_paragraph()
    no_space(gap)
    gap.paragraph_format.space_after = Pt(6)

def dark_box(doc, text):
    """Dark quote/highlight box."""
    p = doc.add_paragraph()
    no_space(p)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(8)
    p.paragraph_format.left_indent  = Inches(0.15)
    para_color(p, '1C1C1C')
    r = p.add_run(f'  {text}')
    set_font(r, size=10.5, color=WHITE, italic=True)

# ══════════════════════════════════════════════════════════════════════════════
# COVER
# ══════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
no_space(p)
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(0)
para_color(p, '1C1C1C')
r = p.add_run('  Super de Alimentos  ·  Trade Marketing')
set_font(r, size=9, color=MGRAY, italic=True)

p2 = doc.add_paragraph()
no_space(p2)
p2.paragraph_format.space_after = Pt(2)
para_color(p2, '1C1C1C')
r2 = p2.add_run('  🌟  OFERTA DE CODIFICACIÓN')
set_font(r2, size=22, bold=True, color=WHITE)

p3 = doc.add_paragraph()
no_space(p3)
p3.paragraph_format.space_after = Pt(0)
para_color(p3, '1C1C1C')
r3 = p3.add_run('  Trululu Nanos — Edición Estelar  |  Ecuador  🇪🇨')
set_font(r3, size=14, bold=False, color=RGBColor(0xAA,0xFF,0x00))

p4 = doc.add_paragraph()
no_space(p4)
p4.paragraph_format.space_after = Pt(4)
para_color(p4, '80CC00')
r4 = p4.add_run('  Estrategia para el Gerente de Trade Marketing  ·  Mayo 2026')
set_font(r4, size=9.5, bold=False, color=WHITE)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ══════════════════════════════════════════════════════════════════════════════
# 1. LECTURA ESTRATÉGICA
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '1. Lectura Estratégica del Contexto')

add_body(doc,
    'Antes de construir cualquier oferta, el contexto entrega tres ventajas competitivas '
    'que deben orientar cada decisión táctica:')

doc.add_paragraph().paragraph_format.space_after = Pt(3)

add_table(doc,
    ['Dato del contexto', 'Implicación táctica'],
    [
        ('Trululu lidera el TOM de la categoría',
         'No hay que vender la marca — ya está instalada. Solo hay que reducir la barrera de entrada a la nueva referencia.'),
        ('80 % del volumen pasa por mayoristas',
         'El mayorista es donde el tendero compra. Activar el punto de venta en mayoristas genera demanda pull hacia el tendero.'),
        ('Un único distribuidor con 785 vendedores llega T2T',
         'El distribuidor es el canal de llegada. Si sus vendedores no están convencidos y equipados, la referencia no llega a la tienda.'),
    ],
    col_widths=[2.5, 4.3])

dark_box(doc,
    '"Edición Estelar es la versión más especial de Trululu Nanos del año. '
    'Tus clientes ya conocen Trululu — esta referencia viene con empaque '
    'coleccionable. Si la tienes, ellos la compran. Si no la tienes, van a otra tienda."')

# ══════════════════════════════════════════════════════════════════════════════
# 2. CONFIGURACIÓN DE LA OFERTA
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '2. Configuración de la Oferta al Tendero')

add_body(doc,
    'La oferta se estructura en tres pilares encadenados. El primero reduce el riesgo de '
    'entrada, el segundo hace que la primera compra sea rentable de inmediato y el tercero '
    'convierte esa primera compra en un hábito de recompra.')

# PILAR 1
add_heading(doc, 'Pilar 1 — Entrada fácil: reduce la barrera de codificación', level=2)

add_body(doc,
    'El principal freno del tendero ante una referencia nueva es el miedo al inventario muerto. '
    'Este pilar elimina ese freno bajando el mínimo de compra y asegurando que el producto '
    'ya tiene demanda instalada en el consumidor.')

add_table(doc,
    ['Elemento', 'Detalle'],
    [
        ('Pack mínimo de codificación',  '2 displays × 24 unidades = 48 unidades'),
        ('Precio de lanzamiento',        '12 % de descuento vs precio regular de lista'),
        ('Vigencia del descuento',       'Solo para la primera compra de la referencia — crea urgencia'),
        ('Exhibidor de sobremesa gratis','Para los primeros 200 tenderos que codifiquen — crea competencia entre tiendas'),
    ],
    col_widths=[2.8, 4.0])

add_body(doc,
    '▸  El exhibidor es clave: el producto queda visible en el mostrador sin que el tendero '
    'tenga que "armar" nada. Reduce la fricción de ejecución.',
    indent=True, color=DARK, size=10)

# PILAR 2
add_heading(doc, 'Pilar 2 — Beneficio inmediato: la primera compra debe ser rentable ya', level=2)

add_body(doc,
    'El tendero evalúa cualquier producto nuevo con una pregunta: ¿en cuánto tiempo recupero '
    'lo que invertí? Este pilar garantiza que la primera compra ya tiene retorno visible.')

add_table(doc,
    ['Elemento', 'Detalle'],
    [
        ('Bonificación 1×12',
         'Por cada 48 unidades compradas → 4 unidades gratis (aumenta el margen por unidad vendida)'),
        ('Sticker de fachada',
         '"Punto Oficial Edición Estelar" — genera orgullo de pertenencia y atrae al consumidor que busca la edición'),
        ('Sorteo mensual',
         'El tendero con más unidades codificadas gana una tablet o televisor — crea competencia positiva entre tiendas'),
        ('Kit de impulso en mostrador',
         'Cenefa + precio sugerido de venta visibles desde el exterior de la tienda'),
    ],
    col_widths=[2.8, 4.0])

# PILAR 3
add_heading(doc, 'Pilar 3 — Acelerador de recompra: convertir la codificación en hábito', level=2)

add_body(doc,
    'Una referencia se consolida cuando el tendero la repone automáticamente. '
    'Este pilar construye ese hábito en las primeras 8 semanas de lanzamiento.')

add_table(doc,
    ['Elemento', 'Detalle'],
    [
        ('Descuento por reposición rápida',
         'Si el tendero repone antes de 15 días → 5 % adicional de descuento en la segunda compra'),
        ('Programa "Estrella Constante"',
         '3 compras consecutivas = el tendero recibe material POP personalizado con el nombre de su negocio'),
        ('Argumento de rotación',
         'El vendedor muestra al tendero datos reales de ventas de Trululu en su zona — evidencia que la demanda ya existe'),
    ],
    col_widths=[2.8, 4.0])

lime_box(doc, '  OFERTA COMPLETA EN UNA LÍNEA',
    ['Pack 48 und + 12 % dcto  →  4 und gratis  →  exhibidor gratis (primeros 200)',
     'Reposición antes de 15 días: +5 % dcto  |  3 compras seguidas: POP personalizado',
     'Sticker de fachada + sorteo mensual de tablet/televisor para el mejor codificador'])

# ══════════════════════════════════════════════════════════════════════════════
# 3. ESTRATEGIA 785 VENDEDORES
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '3. Estrategia para Llegar a los 785 Vendedores')

add_body(doc,
    'Los 785 vendedores del distribuidor no son el cliente final — son el canal de llegada. '
    'La estrategia opera en cascada: primero se convence a la cúpula, luego se capacita '
    'y equipa a cada vendedor, y finalmente se hace seguimiento para que ninguno abandone '
    'la campaña en el camino.')

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# Steps
steps = [
    ('PASO 1  —  SEMANA 0',
     'Conquista la cúpula del distribuidor',
     'Si el gerente de ventas del distribuidor no está convencido, los 785 vendedores tampoco lo estarán. '
     'Este es el paso más crítico y el que más se subestima.',
     [
        'Agendar reunión con Gerencia y Jefatura de Ventas del distribuidor (el ejecutivo comercial Ecuador coordina)',
        'Presentar: la oferta completa, el margen real del tendero, el plan de bono para sus vendedores '
         'y el argumento de TOM (Trululu ya está en la mente del consumidor — solo hay que capitalizar)',
        'Entregar el Kit de Lanzamiento del Distribuidor: muestra del producto, calculadora de margen '
         'para el tendero, y el plan de comunicación con los vendedores',
        'Acordar metas de codificación por zona y vendedor con el distribuidor — el distribuidor '
         'se convierte en aliado activo, no en intermediario pasivo',
     ]),
    ('PASO 2  —  SEMANA 1',
     'Capacitación masiva a los 785 vendedores',
     'Un vendedor que no sabe qué vende no vende. Uno que sí sabe pero no tiene el argumento '
     'correcto tampoco convence al tendero.',
     [
        'Dividir a los 785 vendedores en grupos de ~100 por zona geográfica',
        'Sesiones de 30 minutos (presencial o virtual por WhatsApp / Meet según zona)',
        'Contenido de la sesión: (1) qué es Edición Estelar y por qué va a rotar, '
         '(2) cómo presentar la oferta al tendero en 60 segundos, '
         '(3) cuánto gana el tendero con la bonificación (mostrar en números concretos), '
         '(4) cómo y cuánto gana el propio vendedor',
     ]),
    ('PASO 3  —  SEMANA 1',
     'Kit del vendedor + incentivo individual escalonado',
     'El vendedor necesita dos cosas: herramientas para convencer al tendero y un interés '
     'propio para esforzarse más allá de su cuota habitual.',
     [
        'Kit físico por vendedor: muestra del producto + flyer A6 con la oferta resumida '
         '+ QR con video corto del producto (30 seg) para mostrar en el mostrador',
        'Bono escalonado: 10–19 tiendas codificadas = Nivel 1 | '
         '20–34 tiendas = Nivel 2 | 35+ tiendas = Nivel 3 + reconocimiento público',
        'El escalamiento es clave: el vendedor que llega a 10 quiere llegar a 20',
     ]),
    ('PASO 4  —  SEMANAS 2–4',
     'Seguimiento semanal, ranking y cierre',
     'Sin seguimiento, el lanzamiento muere en la semana 1. El ranking visible '
     'genera competencia positiva entre vendedores y zonas.',
     [
        'Dashboard compartido con el distribuidor: codificaciones por vendedor / ruta / ciudad — '
         'actualizado cada semana',
        'Ranking semanal enviado por WhatsApp al grupo de ventas del distribuidor '
         '(Top 10 visible para todos — la transparencia activa la competencia)',
        'El ejecutivo comercial de Ecuador hace seguimiento con la jefatura del distribuidor '
         'dos veces por semana durante el período de lanzamiento',
        'Cierre semana 4: entrega de bonos + reconocimiento del Top 10 en reunión o video '
         'grupal con el equipo de Super — el cierre público refuerza la cultura de logro',
     ]),
]

for (tag, title, rationale, bullets) in steps:
    # Step header
    ph = doc.add_paragraph()
    no_space(ph)
    ph.paragraph_format.space_before = Pt(10)
    ph.paragraph_format.space_after  = Pt(0)
    para_color(ph, '1C1C1C')
    rt = ph.add_run(f'  {tag}')
    set_font(rt, size=9, bold=True, color=MGRAY)

    pt = doc.add_paragraph()
    no_space(pt)
    pt.paragraph_format.space_after = Pt(2)
    para_color(pt, '80CC00')
    rtt = pt.add_run(f'  {title}')
    set_font(rtt, size=12, bold=True, color=WHITE)

    # Rationale
    pr = doc.add_paragraph()
    no_space(pr)
    pr.paragraph_format.space_after  = Pt(4)
    pr.paragraph_format.left_indent  = Inches(0.1)
    para_color(pr, 'F0FFCC')
    rr = pr.add_run(f'  ¿Por qué? {rationale}')
    set_font(rr, size=9.5, color=DARK, italic=True)

    for b in bullets:
        add_bullet(doc, b)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ══════════════════════════════════════════════════════════════════════════════
# 4. KPIS
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '4. Indicadores de Éxito — Semana 4')

add_table(doc,
    ['Indicador', 'Meta semana 4', 'Responsable'],
    [
        ('Tiendas codificadas / universo atendido', '≥ 60 %', 'Ejecutivo Comercial + Distribuidor'),
        ('Vendedores activos en la campaña', '≥ 80 % de los 785', 'Jefatura de Ventas Distribuidor'),
        ('Unidades vendidas Edición Estelar', 'Según forecast de lanzamiento', 'Trade Marketing'),
        ('Tenderos con recompra antes de día 15', '≥ 30 % de codificadores', 'Vendedor + Ejecutivo'),
        ('Exhibidores instalados', '200 (agotamiento del incentivo)', 'Mercaderistas + Distribuidor'),
    ],
    col_widths=[2.8, 1.8, 2.2],
    header_bg='80CC00')

# ══════════════════════════════════════════════════════════════════════════════
# 5. RESUMEN EJECUTIVO
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '5. Resumen Ejecutivo')

lime_box(doc, '  LA OFERTA',
    ['Entrada: 48 und (2 displays) + 12 % dcto + exhibidor gratis (primeros 200)',
     'Beneficio inmediato: bonificación 1×12 + sticker fachada + sorteo mensual',
     'Recompra: +5 % dcto si repone antes de 15 días  |  "Estrella Constante" a los 3 pedidos'])

lime_box(doc, '  LA LLEGADA A 785 VENDEDORES',
    ['Semana 0: Cúpula del distribuidor convencida con kit + metas acordadas',
     'Semana 1: Capacitación en grupos de 100 (30 min) + kit por vendedor + bono escalonado',
     'Semanas 2–4: Dashboard semanal + ranking WhatsApp + cierre con reconocimiento público'])

dark_box(doc,
    'La ventaja más grande que tiene esta campaña es que el trabajo de construcción '
    'de marca ya está hecho. Trululu es el líder de TOM. '
    'El único trabajo pendiente es capitalizar esa confianza '
    'con una referencia nueva que el tendero sienta como una oportunidad, no como un riesgo.')

# Footer note
p_foot = doc.add_paragraph()
no_space(p_foot)
p_foot.paragraph_format.space_before = Pt(16)
r_foot = p_foot.add_run(
    'Super de Alimentos  ·  Assessment Analista Trade Marketing  ·  '
    'Mariana Montejo  ·  Mayo 2026')
set_font(r_foot, size=8.5, color=MGRAY, italic=True)
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── Save ──────────────────────────────────────────────────────────────────────
out = '/home/user/ULTRA-/Caso3_Oferta_Edicion_Estelar.docx'
doc.save(out)
print(f'✅  Guardado: {out}')
